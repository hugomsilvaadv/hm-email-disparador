import base64
import csv
import html
import hmac
import hashlib
import io
import json
import os
import re
from datetime import datetime, date, timezone
from functools import wraps
from urllib import error as urllib_error
from urllib import request as urllib_request

from flask import Flask, flash, redirect, render_template, request, session, url_for, Response
from flask_sqlalchemy import SQLAlchemy
from openpyxl import load_workbook

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "change-this-secret-before-production")
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["MAX_CONTENT_LENGTH"] = 12 * 1024 * 1024
app.config["SESSION_COOKIE_SECURE"] = True
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"

db_url = os.environ.get("DATABASE_URL", "sqlite:///street_mall.db")
if db_url.startswith("postgres://"):
    db_url = db_url.replace("postgres://", "postgresql+psycopg://", 1)
elif db_url.startswith("postgresql://") and "+psycopg" not in db_url:
    db_url = db_url.replace("postgresql://", "postgresql+psycopg://", 1)
app.config["SQLALCHEMY_DATABASE_URI"] = db_url
db = SQLAlchemy(app)

PROPERTY_STATUS = [
    "Mapeado",
    "Validar anúncio",
    "Contato pendente",
    "Contato realizado",
    "Documentação solicitada",
    "Due diligence",
    "Shortlist",
    "Descartado",
]
LISTING_STATUS = ["A confirmar", "Ativo", "Indisponível", "Vendido/locado", "Duplicado"]
PRIORITY_OPTIONS = ["Alta", "Média", "Baixa"]
CONTACT_STATUS = [
    "Não contatado",
    "Contato localizado",
    "E-mail enviado",
    "Respondeu",
    "Em negociação",
    "Sem interesse",
    "Não contatar",
]
CONTACT_ROLES = ["Proprietário", "Corretor", "Imobiliária", "Captador", "Outro"]

DEFAULT_SUBJECT = "Consulta sobre imóvel comercial em {city}"
DEFAULT_BODY = """Olá, {contact_name}.

Meu nome é Hugo Mendes. Estamos avaliando oportunidades imobiliárias para desenvolvimento de projetos de street mall / open mall em {city}.

Identificamos o imóvel localizado em {address}{area_line} e gostaríamos de confirmar algumas informações antes de avançarmos na análise:

• disponibilidade atual do imóvel;
• valor e condições de negociação;
• matrícula e titularidade;
• testada e acessos;
• zoneamento / uso comercial;
• levantamento ou planta disponível;
• possibilidade de venda, locação de longo prazo, parceria, BTS ou ground lease, quando aplicável.

Caso o imóvel permaneça disponível, podemos seguir por este e-mail ou WhatsApp para uma conversa objetiva.

Atenciosamente,
Hugo Mendes
"""


class Property(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    city = db.Column(db.String(120), nullable=False, index=True)
    state = db.Column(db.String(8), index=True)
    address = db.Column(db.String(500), nullable=False)
    neighborhood = db.Column(db.String(180))
    area_m2 = db.Column(db.Float)
    asking_price = db.Column(db.Float)
    rent_price = db.Column(db.Float)
    price_per_m2 = db.Column(db.Float)
    listing_type = db.Column(db.String(180))
    suggested_format = db.Column(db.String(180))
    priority = db.Column(db.String(40), default="Média", index=True)
    preliminary_score = db.Column(db.Float, index=True)
    location_score = db.Column(db.Float)
    access_score = db.Column(db.Float)
    geometry_score = db.Column(db.Float)
    economics_score = db.Column(db.Float)
    catchment_score = db.Column(db.Float)
    risk_score = db.Column(db.Float)
    rationale = db.Column(db.Text)
    risks = db.Column(db.Text)
    next_action = db.Column(db.Text)
    source_url = db.Column(db.String(900))
    research_date = db.Column(db.Date)
    status = db.Column(db.String(60), default="Mapeado", index=True)
    listing_status = db.Column(db.String(60), default="A confirmar", index=True)
    frontage_m = db.Column(db.Float)
    corner_lot = db.Column(db.Boolean, default=False)
    zoning = db.Column(db.String(240))
    access_notes = db.Column(db.Text)
    parking_notes = db.Column(db.Text)
    infrastructure_notes = db.Column(db.Text)
    owner_name = db.Column(db.String(240))
    registry_number = db.Column(db.String(160))
    latitude = db.Column(db.Float)
    longitude = db.Column(db.Float)
    verified_at = db.Column(db.DateTime(timezone=True))
    notes = db.Column(db.Text)
    created_at = db.Column(db.DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    updated_at = db.Column(db.DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), onupdate=lambda: datetime.now(timezone.utc))

    contacts = db.relationship("PropertyContact", backref="property", cascade="all, delete-orphan", lazy=True)

    @property
    def score(self):
        components = [
            self.location_score,
            self.access_score,
            self.geometry_score,
            self.economics_score,
            self.catchment_score,
        ]
        filled = [v for v in components if v is not None]
        if filled:
            base = sum(filled) / len(filled)
            if self.risk_score is not None:
                base = max(0, base - max(0, self.risk_score - 5) * 0.15)
            return round(base, 1)
        return self.preliminary_score

    @property
    def display_city(self):
        return f"{self.city}/{self.state}" if self.state and "/" not in self.city else self.city

    @property
    def has_contact(self):
        return any(c.email or c.phone or c.whatsapp for c in self.contacts)


class PropertyContact(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    property_id = db.Column(db.Integer, db.ForeignKey("property.id"), nullable=False, index=True)
    name = db.Column(db.String(240))
    role = db.Column(db.String(80))
    company = db.Column(db.String(240))
    email = db.Column(db.String(240), index=True)
    phone = db.Column(db.String(80))
    whatsapp = db.Column(db.String(80))
    source_url = db.Column(db.String(900))
    status = db.Column(db.String(80), default="Não contatado", index=True)
    last_contact = db.Column(db.DateTime(timezone=True))
    next_follow_up = db.Column(db.Date)
    notes = db.Column(db.Text)
    created_at = db.Column(db.DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class Interaction(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    property_id = db.Column(db.Integer, db.ForeignKey("property.id"), nullable=False, index=True)
    contact_id = db.Column(db.Integer, db.ForeignKey("property_contact.id"), nullable=True, index=True)
    occurred_at = db.Column(db.DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), index=True)
    channel = db.Column(db.String(60))
    summary = db.Column(db.Text, nullable=False)
    next_action = db.Column(db.Text)
    follow_up_date = db.Column(db.Date)
    contact = db.relationship("PropertyContact")
    property = db.relationship("Property")


class EmailLog(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    property_id = db.Column(db.Integer, db.ForeignKey("property.id"), nullable=False, index=True)
    contact_id = db.Column(db.Integer, db.ForeignKey("property_contact.id"), nullable=False, index=True)
    sent_at = db.Column(db.DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), index=True)
    recipient = db.Column(db.String(240), nullable=False)
    subject = db.Column(db.String(400), nullable=False)
    body = db.Column(db.Text)
    result = db.Column(db.String(40), nullable=False)
    detail = db.Column(db.Text)
    contact = db.relationship("PropertyContact")
    property = db.relationship("Property")


def env_bool(name, default=False):
    return os.environ.get(name, str(default)).strip().lower() in {"1", "true", "yes", "on"}


def login_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login", next=request.path))
        return fn(*args, **kwargs)
    return wrapper


def clean_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def to_float(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value).strip()
    raw = re.sub(r"[^\d,.\-]", "", raw)
    if not raw:
        return None
    if "," in raw and "." in raw:
        if raw.rfind(",") > raw.rfind("."):
            raw = raw.replace(".", "").replace(",", ".")
        else:
            raw = raw.replace(",", "")
    elif "," in raw:
        raw = raw.replace(".", "").replace(",", ".")
    try:
        return float(raw)
    except ValueError:
        return None


def to_date(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    return None


def split_city_state(raw):
    text = clean_text(raw)
    if "/" in text:
        city, state = text.rsplit("/", 1)
        return city.strip(), state.strip().upper()
    return text, ""


def field(row, *names):
    for name in names:
        if name in row and row[name] not in (None, ""):
            return row[name]
    lowered = {clean_text(k).lower(): v for k, v in row.items()}
    for name in names:
        if name.lower() in lowered and lowered[name.lower()] not in (None, ""):
            return lowered[name.lower()]
    return None


def property_payload_from_row(row):
    city_raw = field(row, "Cidade", "city")
    address = clean_text(field(row, "Endereço / eixo", "Endereco / eixo", "Endereço", "address"))
    if not city_raw or not address:
        return None
    city, state = split_city_state(city_raw)
    area = to_float(field(row, "Área (m²)", "Area (m2)", "Área", "area_m2"))
    asking = to_float(field(row, "Preço pedido (R$)", "Preco pedido (R$)", "asking_price"))
    price_m2 = to_float(field(row, "R$/m²", "R$/m2", "price_per_m2"))
    if price_m2 is None and area and asking:
        price_m2 = asking / area
    return {
        "city": city,
        "state": state,
        "address": address,
        "area_m2": area,
        "asking_price": asking,
        "price_per_m2": price_m2,
        "listing_type": clean_text(field(row, "Tipo / anúncio", "Tipo / anuncio", "listing_type")),
        "suggested_format": clean_text(field(row, "Formato sugerido", "suggested_format")),
        "priority": clean_text(field(row, "Prioridade", "priority")) or "Média",
        "preliminary_score": to_float(field(row, "Nota preliminar (0-10)", "Nota preliminar", "score")),
        "rationale": clean_text(field(row, "Por que interessa", "rationale")),
        "risks": clean_text(field(row, "Riscos / validar", "risks")),
        "next_action": clean_text(field(row, "Próxima ação", "Proxima acao", "next_action")),
        "source_url": clean_text(field(row, "Fonte / anúncio", "Fonte / anuncio", "source_url")),
        "research_date": to_date(field(row, "Data pesquisa", "research_date")),
    }


def upsert_property(payload):
    query = None
    if payload.get("source_url"):
        query = Property.query.filter(Property.source_url == payload["source_url"]).first()
    if not query:
        query = Property.query.filter(
            db.func.lower(Property.city) == payload["city"].lower(),
            db.func.lower(Property.address) == payload["address"].lower(),
        ).first()
    if query:
        for key, value in payload.items():
            if value not in (None, ""):
                setattr(query, key, value)
        return "updated"
    db.session.add(Property(**payload))
    return "created"


def import_rows(rows):
    created = updated = skipped = 0
    for row in rows:
        payload = property_payload_from_row(row)
        if not payload:
            skipped += 1
            continue
        result = upsert_property(payload)
        if result == "created":
            created += 1
        else:
            updated += 1
    db.session.commit()
    return created, updated, skipped


def read_xlsx_rows(file_storage):
    wb = load_workbook(file_storage.stream, data_only=True, read_only=True)
    ws = wb["Pipeline Terrenos"] if "Pipeline Terrenos" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    headers = None
    for idx, row in enumerate(rows[:15]):
        normalized = [clean_text(v) for v in row]
        if "Cidade" in normalized and any("Endereço" in x or "Endereco" in x for x in normalized):
            header_idx = idx
            headers = normalized
            break
    if header_idx is None:
        raise ValueError("Não encontrei a linha de cabeçalho da planilha. Use a aba 'Pipeline Terrenos'.")
    result = []
    for values in rows[header_idx + 1:]:
        if not any(v not in (None, "") for v in values):
            continue
        row = {headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}
        result.append(row)
    return result


def resend_ready():
    return bool(os.environ.get("RESEND_API_KEY", "").strip() and os.environ.get("FROM_EMAIL", "").strip())


def send_enabled():
    return env_bool("SEND_ENABLED", False)


def text_to_html(text):
    safe = html.escape(text)
    return "".join(
        f"<p style='margin:0 0 14px 0;line-height:1.55'>{block.replace(chr(10), '<br>')}</p>"
        for block in safe.split("\n\n")
    )


def send_via_resend(recipient, subject, text_body):
    api_key = os.environ.get("RESEND_API_KEY", "").strip()
    from_email = os.environ.get("FROM_EMAIL", "").strip()
    reply_to = os.environ.get("REPLY_TO", from_email).strip()
    from_name = os.environ.get("FROM_NAME", "Hugo Mendes | Street Mall").strip()
    if not api_key or not from_email:
        raise RuntimeError("Configure RESEND_API_KEY e FROM_EMAIL no Railway.")
    payload = {
        "from": f"{from_name} <{from_email}>",
        "to": [recipient],
        "subject": subject,
        "text": text_body,
        "html": (
            "<!doctype html><html><body style='font-family:Arial,Helvetica,sans-serif;color:#18352f'>"
            "<div style='max-width:680px;margin:auto;padding:24px'>"
            "<div style='border-top:5px solid #b89245;padding-top:18px'>"
            + text_to_html(text_body)
            + "</div></div></body></html>"
        ),
        "reply_to": reply_to,
    }
    req = urllib_request.Request(
        "https://api.resend.com/emails",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib_request.urlopen(req, timeout=20) as response:
            raw = response.read().decode("utf-8")
            if response.status < 200 or response.status >= 300:
                raise RuntimeError(f"Resend retornou HTTP {response.status}.")
            return json.loads(raw) if raw else {}
    except urllib_error.HTTPError as exc:
        try:
            detail = exc.read().decode("utf-8")
        except Exception:
            detail = str(exc)
        raise RuntimeError(f"Resend HTTP {exc.code}: {detail[:700]}") from exc
    except urllib_error.URLError as exc:
        raise RuntimeError(f"Falha de conexão com a Resend: {exc.reason}") from exc


def default_message(contact):
    prop = contact.property
    area_line = f", com aproximadamente {prop.area_m2:,.0f} m²".replace(",", ".") if prop.area_m2 else ""
    name = contact.name or "tudo bem?"
    return (
        DEFAULT_SUBJECT.format(city=prop.display_city),
        DEFAULT_BODY.format(
            contact_name=name,
            city=prop.display_city,
            address=prop.address,
            area_line=area_line,
        ),
    )


def get_property_form_data(prop=None):
    source = request.form
    prop = prop or Property()
    prop.city = clean_text(source.get("city")) or prop.city
    prop.state = clean_text(source.get("state")).upper() or prop.state
    prop.address = clean_text(source.get("address")) or prop.address
    prop.neighborhood = clean_text(source.get("neighborhood"))
    prop.area_m2 = to_float(source.get("area_m2"))
    prop.asking_price = to_float(source.get("asking_price"))
    prop.rent_price = to_float(source.get("rent_price"))
    prop.price_per_m2 = to_float(source.get("price_per_m2"))
    if prop.price_per_m2 is None and prop.area_m2 and prop.asking_price:
        prop.price_per_m2 = prop.asking_price / prop.area_m2
    prop.listing_type = clean_text(source.get("listing_type"))
    prop.suggested_format = clean_text(source.get("suggested_format"))
    prop.priority = clean_text(source.get("priority")) or "Média"
    prop.preliminary_score = to_float(source.get("preliminary_score"))
    prop.location_score = to_float(source.get("location_score"))
    prop.access_score = to_float(source.get("access_score"))
    prop.geometry_score = to_float(source.get("geometry_score"))
    prop.economics_score = to_float(source.get("economics_score"))
    prop.catchment_score = to_float(source.get("catchment_score"))
    prop.risk_score = to_float(source.get("risk_score"))
    prop.rationale = clean_text(source.get("rationale"))
    prop.risks = clean_text(source.get("risks"))
    prop.next_action = clean_text(source.get("next_action"))
    prop.source_url = clean_text(source.get("source_url"))
    prop.research_date = to_date(source.get("research_date"))
    prop.status = clean_text(source.get("status")) or "Mapeado"
    prop.listing_status = clean_text(source.get("listing_status")) or "A confirmar"
    prop.frontage_m = to_float(source.get("frontage_m"))
    prop.corner_lot = source.get("corner_lot") == "on"
    prop.zoning = clean_text(source.get("zoning"))
    prop.access_notes = clean_text(source.get("access_notes"))
    prop.parking_notes = clean_text(source.get("parking_notes"))
    prop.infrastructure_notes = clean_text(source.get("infrastructure_notes"))
    prop.owner_name = clean_text(source.get("owner_name"))
    prop.registry_number = clean_text(source.get("registry_number"))
    prop.latitude = to_float(source.get("latitude"))
    prop.longitude = to_float(source.get("longitude"))
    prop.notes = clean_text(source.get("notes"))
    return prop


@app.get("/health")
def health():
    return {"ok": True, "service": "street-mall-crm"}, 200


@app.route("/login", methods=["GET", "POST"])
def login():
    admin_user = os.environ.get("ADMIN_USER", "hugo")
    login_hash = os.environ.get("LOGIN_SHA256", "").strip().lower()
    legacy_password = os.environ.get("ADMIN_PASSWORD", "")
    if request.method == "POST":
        supplied_user = request.form.get("username", "")
        supplied_password = request.form.get("password", "")
        valid_user = hmac.compare_digest(supplied_user, admin_user)
        supplied_hash = hashlib.sha256(supplied_password.encode("utf-8")).hexdigest()
        if login_hash:
            valid_password = hmac.compare_digest(supplied_hash, login_hash)
        else:
            valid_password = bool(legacy_password) and hmac.compare_digest(supplied_password, legacy_password)
        if valid_user and valid_password:
            session["logged_in"] = True
            return redirect(request.args.get("next") or url_for("dashboard"))
        configured = bool(login_hash or legacy_password)
        flash("Usuário ou senha inválidos." if configured else "Credencial de login ainda não foi configurada.", "danger")
    return render_template("login.html", admin_user=admin_user)


@app.get("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.get("/")
@login_required
def dashboard():
    today = date.today()
    total = Property.query.count()
    high = Property.query.filter(Property.priority == "Alta").count()
    shortlist = Property.query.filter(Property.status == "Shortlist").count()
    with_contact = Property.query.join(PropertyContact).distinct().count()
    city_counts = (
        db.session.query(Property.city, Property.state, db.func.count(Property.id))
        .group_by(Property.city, Property.state)
        .order_by(db.func.count(Property.id).desc())
        .all()
    )
    top_properties = Property.query.order_by(Property.preliminary_score.desc().nullslast(), Property.priority.asc()).limit(8).all()
    followups = (
        PropertyContact.query.filter(PropertyContact.next_follow_up.isnot(None), PropertyContact.next_follow_up <= today)
        .order_by(PropertyContact.next_follow_up.asc())
        .limit(10).all()
    )
    return render_template(
        "dashboard.html",
        total=total,
        high=high,
        shortlist=shortlist,
        with_contact=with_contact,
        city_counts=city_counts,
        top_properties=top_properties,
        followups=followups,
        resend_ready=resend_ready(),
        send_enabled=send_enabled(),
    )


@app.route("/properties", methods=["GET", "POST"])
@login_required
def properties():
    if request.method == "POST":
        file = request.files.get("file")
        if not file or not file.filename:
            flash("Selecione a planilha ou CSV.", "warning")
            return redirect(url_for("properties"))
        try:
            filename = file.filename.lower()
            if filename.endswith(".xlsx"):
                rows = read_xlsx_rows(file)
            elif filename.endswith(".csv"):
                raw = file.read().decode("utf-8-sig")
                rows = csv.DictReader(io.StringIO(raw))
            else:
                raise ValueError("Formato não suportado. Use XLSX ou CSV.")
            created, updated, skipped = import_rows(rows)
            flash(f"Importação concluída: {created} novo(s), {updated} atualizado(s), {skipped} ignorado(s).", "success")
        except Exception as exc:
            flash(f"Falha na importação: {exc}", "danger")
        return redirect(url_for("properties"))

    q = Property.query
    search = request.args.get("q", "").strip()
    city = request.args.get("city", "").strip()
    priority = request.args.get("priority", "").strip()
    status = request.args.get("status", "").strip()
    listing_status = request.args.get("listing_status", "").strip()
    if search:
        term = f"%{search}%"
        q = q.filter(db.or_(Property.address.ilike(term), Property.city.ilike(term), Property.neighborhood.ilike(term)))
    if city:
        q = q.filter(Property.city == city)
    if priority:
        q = q.filter(Property.priority == priority)
    if status:
        q = q.filter(Property.status == status)
    if listing_status:
        q = q.filter(Property.listing_status == listing_status)
    rows = q.order_by(Property.preliminary_score.desc().nullslast(), Property.priority.asc(), Property.city.asc()).all()
    cities = [r[0] for r in db.session.query(Property.city).distinct().order_by(Property.city).all() if r[0]]
    return render_template(
        "properties.html",
        properties=rows,
        cities=cities,
        priority_options=PRIORITY_OPTIONS,
        status_options=PROPERTY_STATUS,
        listing_status_options=LISTING_STATUS,
    )


@app.route("/properties/new", methods=["GET", "POST"])
@login_required
def property_new():
    if request.method == "POST":
        prop = get_property_form_data()
        if not prop.city or not prop.address:
            flash("Cidade e endereço/eixo são obrigatórios.", "danger")
        else:
            db.session.add(prop)
            db.session.commit()
            flash("Imóvel cadastrado.", "success")
            return redirect(url_for("property_detail", property_id=prop.id))
    return render_template("property_form.html", prop=None, priority_options=PRIORITY_OPTIONS, status_options=PROPERTY_STATUS, listing_status_options=LISTING_STATUS)


@app.route("/properties/<int:property_id>/edit", methods=["GET", "POST"])
@login_required
def property_edit(property_id):
    prop = db.get_or_404(Property, property_id)
    if request.method == "POST":
        prop = get_property_form_data(prop)
        if not prop.city or not prop.address:
            flash("Cidade e endereço/eixo são obrigatórios.", "danger")
        else:
            db.session.commit()
            flash("Ficha atualizada.", "success")
            return redirect(url_for("property_detail", property_id=prop.id))
    return render_template("property_form.html", prop=prop, priority_options=PRIORITY_OPTIONS, status_options=PROPERTY_STATUS, listing_status_options=LISTING_STATUS)


@app.get("/properties/<int:property_id>")
@login_required
def property_detail(property_id):
    prop = db.get_or_404(Property, property_id)
    interactions = Interaction.query.filter_by(property_id=prop.id).order_by(Interaction.occurred_at.desc()).all()
    logs = EmailLog.query.filter_by(property_id=prop.id).order_by(EmailLog.sent_at.desc()).limit(20).all()
    return render_template(
        "property_detail.html",
        prop=prop,
        interactions=interactions,
        logs=logs,
        contact_status_options=CONTACT_STATUS,
        contact_roles=CONTACT_ROLES,
    )


@app.post("/properties/<int:property_id>/verify")
@login_required
def property_verify(property_id):
    prop = db.get_or_404(Property, property_id)
    prop.verified_at = datetime.now(timezone.utc)
    if prop.listing_status == "A confirmar":
        prop.listing_status = "Ativo"
    db.session.commit()
    flash("Anúncio marcado como verificado agora.", "success")
    return redirect(url_for("property_detail", property_id=prop.id))


@app.post("/properties/<int:property_id>/contacts")
@login_required
def contact_add(property_id):
    prop = db.get_or_404(Property, property_id)
    contact = PropertyContact(
        property_id=prop.id,
        name=clean_text(request.form.get("name")),
        role=clean_text(request.form.get("role")),
        company=clean_text(request.form.get("company")),
        email=clean_text(request.form.get("email")).lower(),
        phone=clean_text(request.form.get("phone")),
        whatsapp=clean_text(request.form.get("whatsapp")),
        source_url=clean_text(request.form.get("source_url")),
        status=clean_text(request.form.get("status")) or "Não contatado",
        next_follow_up=to_date(request.form.get("next_follow_up")),
        notes=clean_text(request.form.get("notes")),
    )
    if not any([contact.name, contact.company, contact.email, contact.phone, contact.whatsapp]):
        flash("Preencha ao menos um dado do contato.", "warning")
        return redirect(url_for("property_detail", property_id=prop.id))
    db.session.add(contact)
    db.session.commit()
    flash("Contato adicionado.", "success")
    return redirect(url_for("property_detail", property_id=prop.id))


@app.post("/properties/<int:property_id>/interactions")
@login_required
def interaction_add(property_id):
    prop = db.get_or_404(Property, property_id)
    summary = clean_text(request.form.get("summary"))
    if not summary:
        flash("Informe o resumo do contato.", "warning")
        return redirect(url_for("property_detail", property_id=prop.id))
    contact_id = request.form.get("contact_id")
    interaction = Interaction(
        property_id=prop.id,
        contact_id=int(contact_id) if contact_id and contact_id.isdigit() else None,
        channel=clean_text(request.form.get("channel")),
        summary=summary,
        next_action=clean_text(request.form.get("next_action")),
        follow_up_date=to_date(request.form.get("follow_up_date")),
    )
    db.session.add(interaction)
    if interaction.contact_id:
        contact = db.session.get(PropertyContact, interaction.contact_id)
        if contact:
            contact.last_contact = datetime.now(timezone.utc)
            contact.next_follow_up = interaction.follow_up_date
            if contact.status == "Não contatado":
                contact.status = "Contato localizado"
    if prop.status in ("Mapeado", "Validar anúncio", "Contato pendente"):
        prop.status = "Contato realizado"
    db.session.commit()
    flash("Interação registrada.", "success")
    return redirect(url_for("property_detail", property_id=prop.id))


@app.get("/contacts")
@login_required
def contacts():
    q = PropertyContact.query.join(Property)
    search = request.args.get("q", "").strip()
    status = request.args.get("status", "").strip()
    city = request.args.get("city", "").strip()
    if search:
        term = f"%{search}%"
        q = q.filter(db.or_(
            PropertyContact.name.ilike(term),
            PropertyContact.company.ilike(term),
            PropertyContact.email.ilike(term),
            Property.address.ilike(term),
        ))
    if status:
        q = q.filter(PropertyContact.status == status)
    if city:
        q = q.filter(Property.city == city)
    rows = q.order_by(Property.city.asc(), PropertyContact.company.asc(), PropertyContact.name.asc()).all()
    cities = [r[0] for r in db.session.query(Property.city).distinct().order_by(Property.city).all() if r[0]]
    return render_template("contacts.html", contacts=rows, cities=cities, contact_status_options=CONTACT_STATUS)


@app.get("/outreach")
@login_required
def outreach():
    rows = (
        PropertyContact.query.join(Property)
        .filter(PropertyContact.email.isnot(None), PropertyContact.email != "")
        .filter(PropertyContact.status != "Não contatar")
        .order_by(PropertyContact.status.asc(), Property.priority.asc(), Property.preliminary_score.desc().nullslast())
        .all()
    )
    return render_template("outreach.html", contacts=rows, send_enabled=send_enabled(), resend_ready=resend_ready())


@app.route("/outreach/<int:contact_id>/preview", methods=["GET", "POST"])
@login_required
def outreach_preview(contact_id):
    contact = db.get_or_404(PropertyContact, contact_id)
    default_subject, default_body = default_message(contact)
    subject = request.form.get("subject", default_subject)
    body = request.form.get("body", default_body)
    if request.method == "POST":
        action = request.form.get("action")
        if action == "send":
            if not contact.email or not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", contact.email):
                flash("Esse contato não possui e-mail válido.", "danger")
            elif not send_enabled():
                flash("Envios continuam bloqueados. Defina SEND_ENABLED=true apenas quando a operação estiver pronta.", "warning")
            elif not resend_ready():
                flash("Configure RESEND_API_KEY e FROM_EMAIL no Railway.", "danger")
            else:
                try:
                    send_via_resend(contact.email, subject, body)
                    contact.status = "E-mail enviado"
                    contact.last_contact = datetime.now(timezone.utc)
                    db.session.add(EmailLog(
                        property_id=contact.property_id,
                        contact_id=contact.id,
                        recipient=contact.email,
                        subject=subject,
                        body=body,
                        result="sent",
                    ))
                    db.session.add(Interaction(
                        property_id=contact.property_id,
                        contact_id=contact.id,
                        channel="E-mail",
                        summary=f"E-mail enviado para {contact.email}: {subject}",
                    ))
                    if contact.property.status in ("Mapeado", "Validar anúncio", "Contato pendente"):
                        contact.property.status = "Contato realizado"
                    db.session.commit()
                    flash("E-mail enviado e histórico atualizado.", "success")
                except Exception as exc:
                    db.session.add(EmailLog(
                        property_id=contact.property_id,
                        contact_id=contact.id,
                        recipient=contact.email or "",
                        subject=subject,
                        body=body,
                        result="error",
                        detail=str(exc)[:1800],
                    ))
                    db.session.commit()
                    flash(f"Falha no envio: {exc}", "danger")
    return render_template(
        "outreach_preview.html",
        contact=contact,
        subject=subject,
        body=body,
        send_enabled=send_enabled(),
        resend_ready=resend_ready(),
    )


@app.get("/export/properties.csv")
@login_required
def export_properties():
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "Cidade", "UF", "Endereço / eixo", "Área (m²)", "Preço pedido (R$)", "R$/m²",
        "Tipo / anúncio", "Formato sugerido", "Prioridade", "Nota", "Status", "Status anúncio",
        "Testada (m)", "Zoneamento", "Por que interessa", "Riscos / validar", "Próxima ação", "Fonte / anúncio"
    ])
    for prop in Property.query.order_by(Property.city, Property.address).all():
        writer.writerow([
            prop.city, prop.state, prop.address, prop.area_m2, prop.asking_price, prop.price_per_m2,
            prop.listing_type, prop.suggested_format, prop.priority, prop.score, prop.status, prop.listing_status,
            prop.frontage_m, prop.zoning, prop.rationale, prop.risks, prop.next_action, prop.source_url
        ])
    return Response(
        output.getvalue(),
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=street_mall_imoveis.csv"},
    )


with app.app_context():
    db.create_all()
    if Property.query.count() == 0:
        seed_path = os.path.join(os.path.dirname(__file__), "seed_properties.csv")
        if os.path.exists(seed_path):
            try:
                with open(seed_path, "r", encoding="utf-8-sig", newline="") as seed_file:
                    created, updated, skipped = import_rows(csv.DictReader(seed_file))
                app.logger.info(
                    "Seed Street Mall carregado: %d novos, %d atualizados, %d ignorados.",
                    created, updated, skipped,
                )
            except Exception as exc:
                app.logger.error("Falha ao carregar seed local de imóveis: %s", exc)

    seed_b64 = os.environ.get("PROPERTIES_SEED_CSV_B64", "").strip()
    if seed_b64 and Property.query.count() == 0:
        try:
            seed_csv = base64.b64decode(seed_b64).decode("utf-8-sig")
            import_rows(csv.DictReader(io.StringIO(seed_csv)))
        except Exception as exc:
            app.logger.error("Falha ao carregar seed de imóveis por variável: %s", exc)

    app.logger.info("Street Mall database ready: %d imóveis.", Property.query.count())


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5000")), debug=False)
